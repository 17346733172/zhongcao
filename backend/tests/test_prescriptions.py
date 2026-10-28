from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import json
import time
import unittest
from unittest.mock import patch
from fastapi import HTTPException
from fastapi.testclient import TestClient
from app.config import Settings
from app.main import create_app
from app.security import hash_password
from app.prescriptions.models import Intent, Prescription, Medicine
from app.prescriptions.rules import check_prescription, canonical
from app.prescriptions.files import read_upload, MAX_BYTES, validate_text
from app.prescriptions.extraction import extract, redact_identifiers


def prescription(**updates):
    value = {"age": 58, "sex": "male", "pregnancy": "no", "diagnosis": "测试", "allergies": "未提供",
             "regimen": "每日1剂，水煎服", "doses": 3,
             "medicines": [{"name": "陈皮", "dose": "9", "unit": "g", "dose_scope": "per_dose", "source": "陈皮 9 g"}]}
    value.update(updates)
    return Prescription.model_validate(value)


class PrescriptionRuleTests(unittest.TestCase):
    def test_processed_aconite_and_pinellia_are_flagged_without_confusing_baifuzi(self):
        rx = prescription(medicines=[{"name": "制附子", "dose": "10", "unit": "g"}, {"name": "法半夏", "dose": "9", "unit": "g"}])
        result = check_prescription(rx)
        self.assertTrue(any(f["code"] == "incompatible_pair" for f in result))
        self.assertNotEqual(canonical("白附子"), canonical("附子"))
        rx.medicines[0].name = "白附子"
        self.assertFalse(any(f["code"] == "incompatible_pair" for f in check_prescription(rx)))

    def test_duplicate_alias_missing_invalid_and_unknown_units(self):
        rx = prescription(age=None, allergies="", medicines=[{"name": "炙甘草", "dose": "NaN", "unit": "钱"}, {"name": "甘草", "dose": "0", "unit": ""}])
        codes = {f["code"] for f in check_prescription(rx)}
        self.assertTrue({"duplicate_herb", "invalid_dose", "unit_review", "missing_dose", "missing_age", "missing_allergies"} <= codes)

    def test_reference_dose_is_not_compared_to_total_unknown_or_child_dose(self):
        rx = prescription(medicines=[{"name": "陈皮", "dose": "999", "unit": "g", "dose_scope": "total"}])
        self.assertFalse(any(f["code"] == "reference_dose" for f in check_prescription(rx)))
        rx.medicines[0].dose_scope = "per_dose"
        self.assertTrue(any(f["code"] == "reference_dose" for f in check_prescription(rx)))
        rx.age = 8
        self.assertFalse(any(f["code"] == "reference_dose" for f in check_prescription(rx)))

    def test_confirmed_total_is_review_not_missing_and_per_dose_clears_scope_findings(self):
        rx = prescription(medicines=[{"name": "党参", "dose": "12", "unit": "g", "dose_scope": "total"}])
        result = check_prescription(rx)
        self.assertFalse(any(f["code"] == "dose_scope" for f in result))
        self.assertTrue(any(f["code"] == "dose_total_review" and f["level"] == "review" for f in result))
        rx.medicines[0].dose_scope = "per_dose"
        result = check_prescription(rx)
        self.assertFalse(any(f["code"] in {"dose_scope", "dose_total_review"} for f in result))

    def test_dose_exponents_and_zero_do_not_create_large_or_nonfinite_totals(self):
        for dose in ["1e999999999", "Infinity", "-1", "0", "9/3"]:
            self.assertTrue(any(f["code"] == "invalid_dose" for f in check_prescription(prescription(medicines=[{"name": "陈皮", "dose": dose, "unit": "g", "dose_scope": "per_dose"}]))))

    def test_grams_and_milligrams_aggregate_for_duplicate_review(self):
        rx = prescription(medicines=[{"name": "陈皮", "dose": "8", "unit": "g", "dose_scope": "per_dose"}, {"name": "陈皮", "dose": "8000", "unit": "mg", "dose_scope": "per_dose"}])
        self.assertTrue(any(f["code"] == "reference_dose" for f in check_prescription(rx)))

    def test_missing_pregnancy_and_unknown_herb_are_not_assumed_safe(self):
        result = check_prescription(prescription(sex="female", pregnancy="unknown", medicines=[{"name": "OCR不明药名", "dose": "9", "unit": "片"}]))
        self.assertTrue({"missing_pregnancy", "unknown_herb", "dose_scope"} <= {f["code"] for f in result})

    def test_unspecified_history_and_conflicting_metadata_are_flagged(self):
        codes = {f["code"] for f in check_prescription(prescription(allergies="未提供", pregnancy="yes"))}
        self.assertIn("missing_allergies", codes)
        self.assertIn("conflicting_metadata", codes)

    def test_extraction_must_match_real_drug_and_numeric_evidence(self):
        settings = Settings(deepseek_api_key="unit-test")
        for medicine in [
            {"name": "陈皮", "dose": "19", "unit": "g", "source": "陈皮 9 g"},
            {"name": "附子", "dose": "10", "unit": "g", "source": "附子 10 g"},
        ]:
            payload = prescription(medicines=[medicine]).model_dump()
            with patch("app.prescriptions.extraction.request_json", return_value=json.dumps(payload)):
                with self.assertRaises(HTTPException) as error:
                    extract(settings, "陈皮 9 g")
                self.assertEqual(error.exception.status_code, 502)

    def test_common_identifiers_are_removed_before_model_calls(self):
        result = redact_identifiers("姓名：虚构姓名\n手机号：13800000000\n陈皮 9 g")
        self.assertNotIn("虚构姓名", result)
        self.assertNotIn("13800000000", result)
        self.assertIn("陈皮 9 g", result)

    def test_redaction_preserves_explicit_age_and_sex_from_the_patient_row(self):
        for text in ["| 患者 | 李某某（化名），男，58 岁 |", "姓名：李某某，性别：男，年龄：58岁\n陈皮 9 g"]:
            with self.subTest(text=text):
                result = redact_identifiers(text)
                self.assertNotIn("李某某", result)
                self.assertIn("性别：男", result)
                self.assertRegex(result, r"年龄：58\s*岁")
        self.assertNotIn("性别", redact_identifiers("姓名：李男"), "a name containing 男 is not evidence of sex")

    def test_negative_dose_from_original_can_be_confirmed_and_flagged(self):
        payload = prescription(medicines=[{"name": "陈皮", "dose": "-1", "unit": "g", "source": "陈皮 -1 g"}]).model_dump()
        with patch("app.prescriptions.extraction.request_json", return_value=json.dumps(payload)):
            result = extract(Settings(deepseek_api_key="unit-test"), "陈皮 -1 g")
        self.assertTrue(any(f["code"] == "invalid_dose" for f in check_prescription(result)))


class PrescriptionFileTests(unittest.TestCase):
    def test_text_encoding_size_empty_and_type_boundaries(self):
        self.assertIn("陈皮", read_upload("陈皮 9 g 每剂".encode(), "处方.txt")["text"])
        for content, name in [(b"", "a.txt"), (b"\xff\xfe", "a.txt"), (b"x" * (MAX_BYTES + 1), "a.txt"), (b"not an image", "a.png"), (b"fake", "a.exe"), (b"fake", "a.docx")]:
            with self.assertRaises(HTTPException): read_upload(content, name)
        with self.assertRaises(HTTPException): validate_text("陈" * 20001)

    def test_docx_tables_keep_document_order(self):
        from docx import Document
        document = Document(); document.add_paragraph("处方开始")
        row = document.add_table(rows=1, cols=2).rows[0]; row.cells[0].text = "陈皮"; row.cells[1].text = "9 g"
        document.add_paragraph("处方结束"); stream = BytesIO(); document.save(stream)
        text = read_upload(stream.getvalue(), "a.docx")["text"]
        self.assertLess(text.index("处方开始"), text.index("陈皮")); self.assertLess(text.index("陈皮"), text.index("处方结束"))

    def test_pdf_corrupt_and_page_limits(self):
        import pypdfium2 as pdfium
        with self.assertRaises(HTTPException): read_upload(b"fake pdf", "a.pdf")
        document = pdfium.PdfDocument.new()
        for _ in range(9): document.new_page(500, 500).close()
        stream = BytesIO(); document.save(stream); document.close()
        with self.assertRaises(HTTPException): read_upload(stream.getvalue(), "a.pdf")

    def test_scanned_pdf_is_rendered_and_uses_ocr(self):
        from PIL import Image
        stream = BytesIO(); Image.new("RGB", (500, 500), "white").save(stream, "PDF")
        with patch("app.prescriptions.files.ocr_image", return_value=("陈皮 9 g 每剂", .8)) as ocr:
            result = read_upload(stream.getvalue(), "scan.pdf")
            self.assertEqual(result["method"], "pdf_ocr")
            self.assertIn("第 1 页", result["text"])
            self.assertTrue(result["notes"])
            ocr.assert_called_once()


class PrescriptionApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.app = create_app(Settings(database=Path(self.temp.name) / "test.db", deepseek_api_key="unit-test"))
        self.client = TestClient(self.app)
        self.client.__enter__()
        with self.app.state.database.connect() as db:
            db.execute("INSERT INTO users(username,display_name,password_hash,role,created_at) VALUES(?,?,?,?,?)", ("rx_viewer", "处方测试", hash_password("unit-test-secret"), "viewer", int(time.time())))

    def tearDown(self):
        self.client.__exit__(None, None, None); self.temp.cleanup()

    def login(self):
        response = self.client.post("/api/auth/login", headers={"X-CSRF-Protection": "1"}, json={"username": "rx_viewer", "password": "unit-test-secret"})
        self.assertEqual(response.status_code, 200)
        self.client.headers["X-CSRF-Protection"] = "1"

    def test_auth_csrf_confirmation_and_private_viewer_access(self):
        self.assertEqual(self.client.get("/api/prescriptions/status").status_code, 401)
        self.login()
        self.assertEqual(self.client.get("/api/prescriptions/status").status_code, 200)
        payload = {"prescription": prescription().model_dump(), "confirmed": False}
        self.assertEqual(self.client.post("/api/prescriptions/check", json=payload).status_code, 422)
        payload["confirmed"] = True
        response = self.client.post("/api/prescriptions/check", json=payload)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["framework"], "LangGraph")
        self.assertEqual(response.json()["steps"], ["check_data", "build_report"])
        self.assertIn("limitations", response.json())
        self.assertEqual(self.client.get("/api/documents").json(), [])
        self.client.headers.pop("X-CSRF-Protection")
        self.assertEqual(self.client.post("/api/prescriptions/check", json=payload).status_code, 403)

    def test_real_graph_prepare_branch_and_clarification_stop(self):
        self.login()
        with patch("app.prescriptions.agent.classify", return_value=Intent(kind="prescription_check", explanation="一份处方")), patch("app.prescriptions.agent.extract", return_value=prescription()) as extracted:
            response = self.client.post("/api/prescriptions/prepare", files={"file": ("rx.txt", "陈皮 9 g 每剂".encode(), "text/plain")})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["steps"], ["read_upload", "recognize_intent", "extract_prescription", "confirm_data"])
            self.assertEqual(response.json()["status"], "needs_confirmation")
            self.assertNotIn("findings", response.json())
            extracted.assert_called_once()
        with patch("app.prescriptions.agent.classify", return_value=Intent(kind="clarify", explanation="有两个版本", multiple_prescriptions=True)), patch("app.prescriptions.agent.extract") as extracted:
            response = self.client.post("/api/prescriptions/prepare", data={"text": "修改前：陈皮9g；修改后：陈皮6g"})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["status"], "needs_clarification")
            extracted.assert_not_called()

    def test_user_per_dose_confirmation_survives_extraction_and_check(self):
        self.login()
        names_and_doses = [("法半夏", "9"), ("干姜", "6"), ("党参", "12"), ("白术", "12"), ("茯苓", "15")]
        extracted = prescription(medicines=[{"name": name, "dose": dose, "unit": "g", "dose_scope": "unknown",
                                             "source": f"{name} {dose} g"} for name, dose in names_and_doses])
        for use_file in [False, True]:
            with self.subTest(use_file=use_file), patch("app.prescriptions.agent.classify", return_value=Intent(kind="prescription_check", explanation="一份处方")), patch("app.prescriptions.agent.extract", return_value=extracted):
                text = "\n".join(m.source for m in extracted.medicines)
                data = {"dose_scope": "per_dose"}
                upload = {"file": ("rx.txt", text.encode(), "text/plain")} if use_file else None
                if not use_file:
                    data["text"] = text
                response = self.client.post("/api/prescriptions/prepare", data=data, files=upload)
                self.assertEqual(response.status_code, 200)
                prepared = response.json()["prescription"]
                self.assertEqual(prepared["dose_scope"], "per_dose")
                self.assertTrue(all(m["dose_scope"] == "per_dose" for m in prepared["medicines"]))
                self.assertEqual([(m["name"], m["dose"], m["unit"]) for m in prepared["medicines"]], [(name, dose, "g") for name, dose in names_and_doses])
                report = self.client.post("/api/prescriptions/check", json={"prescription": prepared, "confirmed": True}).json()
                self.assertEqual(report["prescription"], prepared)
                self.assertFalse(any(f["code"] == "dose_scope" for f in report["findings"]))
                self.assertTrue(any(f["code"] == "missing_allergies" for f in report["findings"]))
        self.assertTrue(all(m.dose_scope == "unknown" for m in extracted.medicines), "normalization must not mutate the model result")

    def test_server_resolves_whole_prescription_choice_before_rules_and_report(self):
        self.login()
        rx = prescription(medicines=[{"name": "陈皮", "dose": "999", "unit": "g", "dose_scope": "unknown"}]).model_dump()
        rx["dose_scope"] = "per_dose"
        report = self.client.post("/api/prescriptions/check", json={"prescription": rx, "confirmed": True}).json()
        self.assertEqual(report["prescription"]["medicines"][0]["dose_scope"], "per_dose")
        self.assertTrue(any(f["code"] == "reference_dose" for f in report["findings"]))
        self.assertFalse(any(f["code"] == "dose_scope" for f in report["findings"]))
        rx["dose_scope"] = "total"
        report = self.client.post("/api/prescriptions/check", json={"prescription": rx, "confirmed": True}).json()
        self.assertEqual(report["prescription"]["medicines"][0]["dose"], "999")
        self.assertEqual(report["prescription"]["medicines"][0]["dose_scope"], "total")
        self.assertFalse(any(f["code"] in {"reference_dose", "dose_scope"} for f in report["findings"]))
        self.assertTrue(any(f["code"] == "dose_total_review" for f in report["findings"]))

    def test_incompatible_pair_conclusion_uses_requested_text_even_with_missing_information(self):
        self.login()
        rx = prescription(dose_scope="per_dose", medicines=[{"name": "制附子", "dose": "10", "unit": "g", "method": "先煎60分钟"}, {"name": "法半夏", "dose": "9", "unit": "g"}])
        response = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True})
        self.assertEqual(response.status_code, 200)
        report = response.json()
        self.assertTrue(any(f["code"] == "missing_allergies" for f in report["findings"]))
        self.assertEqual(report["conclusion"]["status"], "not_passed")
        self.assertEqual(report["conclusion"]["result"], "不通过。")
        self.assertEqual(report["conclusion"]["detail"], "处方存在“十八反”配伍禁忌，建议暂缓调配，联系处方医师确认或修改。")
        self.assertIn("incompatible_pair", report["conclusion"]["finding_codes"])

    def test_conclusion_updates_after_edit_and_does_not_invent_pharmacist_approval(self):
        self.login()
        rx = prescription(allergies="无", dose_scope="per_dose", medicines=[{"name": "陈皮", "dose": "9", "unit": "g", "method": "同煎"}])
        def report():
            return self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
        value = report()
        self.assertEqual(value["findings"], [])
        self.assertEqual(value["conclusion"]["status"], "passed")
        self.assertEqual(value["conclusion"]["result"], "通过。")
        rx.medicines[0].dose = "-1"
        value = report()
        self.assertEqual(value["conclusion"]["status"], "not_passed")
        self.assertNotIn("十八反", value["conclusion"]["detail"])
        rx.medicines[0].dose = "9"
        self.assertEqual(report()["conclusion"]["status"], "passed")
        rx.allergies = ""
        self.assertEqual(report()["conclusion"]["status"], "needs_information")
        rx.allergies = "无"
        rx.medicines[0].dose_scope = "total"
        rx.dose_scope = "unknown"
        self.assertEqual(report()["conclusion"]["status"], "not_passed")

    def test_each_registered_incompatibility_family_blocks_review_but_aconite_alone_does_not(self):
        self.login()
        for names in [("甘草", "甘遂"), ("藜芦", "党参"), ("法半夏", "制附子")]:
            with self.subTest(names=names):
                rx = prescription(dose_scope="per_dose", medicines=[{"name": name, "dose": "3", "unit": "g"} for name in names])
                report = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
                self.assertEqual(report["conclusion"]["status"], "not_passed")
                self.assertIn("十八反", report["conclusion"]["detail"])
        rx = prescription(allergies="无", dose_scope="per_dose", medicines=[{"name": "制附子", "dose": "10", "unit": "g", "method": "先煎60分钟"}])
        report = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
        self.assertEqual(report["conclusion"]["status"], "passed")
        self.assertNotIn("十八反", report["conclusion"]["detail"])

    def test_five_section_report_uses_current_names_doses_methods_and_requested_conclusion(self):
        self.login()
        rx = prescription(allergies="无", pregnancy="no", dose_scope="per_dose", medicines=[
            {"name": "制附子", "dose": "10", "unit": "g", "method": "先煎60分钟"},
            {"name": "法半夏", "dose": "9", "unit": "g", "method": "同煎"},
            {"name": "砂仁", "dose": "6", "unit": "g", "method": "后下，打碎"},
            {"name": "炙甘草", "dose": "6", "unit": "g", "method": "同煎"},
        ])
        report = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
        sections = report["sections"]
        self.assertEqual([s["number"] for s in sections], ["3.1", "3.2", "3.3", "3.4", "3.5"])
        self.assertEqual(sections[0]["paragraphs"][0]["lead"], "不通过。")
        self.assertEqual(sections[0]["paragraphs"][0]["text"], report["conclusion"]["detail"])
        self.assertEqual(sections[1]["title"], "主要问题：制附子与法半夏同方")
        methods = json.dumps(sections[2], ensure_ascii=False)
        for value in ["制附子 10 g", "法半夏 9 g", "先煎60分钟", "后下，打碎"]:
            self.assertIn(value, methods)
        self.assertIn("低钾", json.dumps(sections[3], ensure_ascii=False))
        self.assertNotIn("未发现“十九畏”", json.dumps(sections[3], ensure_ascii=False))
        self.assertTrue(sections[4]["ordered"])
        self.assertTrue(sections[4]["items"][0]["children"])

    def test_narrative_changes_after_edit_and_does_not_copy_example_drugs_or_results(self):
        self.login()
        for name, dose in [("陈皮", "6"), ("桂枝", "3")]:
            with self.subTest(name=name):
                rx = prescription(allergies="无", dose_scope="per_dose", medicines=[{"name": name, "dose": dose, "unit": "g", "method": "同煎"}])
                report = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
                body = json.dumps(report["sections"], ensure_ascii=False)
                self.assertNotIn("制附子与法半夏同方", body)
                self.assertNotIn("制附子 10 g", body)
                self.assertNotIn("首选：", body)
                self.assertNotEqual(report["sections"][0]["paragraphs"][0]["lead"], "不通过。")

    def test_narrative_does_not_compare_total_child_or_unknown_units_as_adult_per_dose(self):
        self.login()
        for changes in [{"age": 8}, {"dose_scope": "total"}, {"medicines": [{"name": "法半夏", "dose": "3", "unit": "片", "dose_scope": "per_dose"}]}]:
            with self.subTest(changes=changes):
                base = {"allergies": "无", "dose_scope": "per_dose", "medicines": [{"name": "法半夏", "dose": "9", "unit": "g"}]}
                base.update(changes)
                report = self.client.post("/api/prescriptions/check", json={"prescription": prescription(**base).model_dump(), "confirmed": True}).json()
                method_section = json.dumps(report["sections"][2], ensure_ascii=False)
                self.assertNotIn("位于公开资料", method_section)
                self.assertNotIn("先煎60分钟", method_section)
                self.assertNotIn("符合要求", method_section)

    def test_narrative_multiple_pairs_and_processing_aliases_remain_distinct(self):
        self.login()
        rx = prescription(allergies="无", dose_scope="per_dose", medicines=[{"name": name, "dose": "3", "unit": "g"} for name in ["白附片", "法半夏", "甘草", "甘遂"]])
        report = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
        self.assertEqual(report["sections"][1]["title"], "主要问题：多组配伍禁忌")
        self.assertIn("白附片与法半夏同方", json.dumps(report["sections"][1], ensure_ascii=False))
        rx.medicines = [Medicine(name="白附子", dose="3", unit="g", dose_scope="per_dose")]
        report = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
        self.assertNotIn("附子用药资料", json.dumps(report["sections"][2], ensure_ascii=False))

    def test_missing_information_lists_only_actual_patient_gaps_in_the_conclusion(self):
        self.login()
        rx = prescription(age=None, allergies="未知", pregnancy="no")
        report = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
        missing = report["missing_information"]
        self.assertEqual([entry["field"] for entry in missing], ["age", "allergies"])
        self.assertEqual([entry["label"] for entry in missing], ["患者年龄", "过敏史"])
        self.assertEqual(report["sections"][0]["paragraphs"][0]["lead"], "待补充信息。")
        self.assertEqual([entry["lead"] for entry in report["sections"][0]["items"]], ["患者年龄：", "过敏史："])
        self.assertNotIn("性别：", json.dumps(report["sections"][0], ensure_ascii=False))

    def test_medicine_missing_information_distinguishes_row_dose_unit_and_scope(self):
        self.login()
        rx = prescription(allergies="无", medicines=[
            {"name": "法半夏", "dose": "9", "unit": "", "dose_scope": "per_dose"},
            {"name": "党参", "dose": "", "unit": "g", "dose_scope": "unknown"},
        ])
        report = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
        self.assertEqual([entry["field"] for entry in report["missing_information"]], ["medicines.0.unit", "medicines.1.dose", "medicines.1.dose_scope"])
        self.assertEqual(report["missing_information"][0]["label"], "第1味药「法半夏」的单位")
        self.assertEqual(report["missing_information"][1]["label"], "第2味药「党参」的剂量")
        self.assertEqual(len(report["sections"][0]["items"]), 3)

    def test_supplemented_information_disappears_without_changing_doses_or_hiding_risks(self):
        self.login()
        rx = prescription(allergies="", dose_scope="per_dose", medicines=[
            {"name": "制附子", "dose": "10", "unit": "g", "method": "先煎60分钟"},
            {"name": "法半夏", "dose": "9", "unit": "g", "method": "同煎"},
        ])
        def report():
            return self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
        first = report()
        self.assertEqual(first["conclusion"]["status"], "not_passed")
        self.assertEqual([entry["field"] for entry in first["missing_information"]], ["allergies"])
        rx.allergies = "无"
        updated = report()
        self.assertEqual(updated["missing_information"], [])
        self.assertEqual(updated["sections"][0]["items"], [])
        self.assertEqual(updated["conclusion"]["status"], "not_passed")
        self.assertEqual([m["dose"] for m in updated["prescription"]["medicines"]], ["10", "9"])

    def test_pregnancy_supplement_is_requested_only_when_the_existing_rule_requires_it(self):
        self.login()
        for sex, fields in [("female", ["pregnancy"]), ("male", [])]:
            with self.subTest(sex=sex):
                rx = prescription(sex=sex, pregnancy="unknown", allergies="无")
                report = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
                self.assertEqual([entry["field"] for entry in report["missing_information"]], fields)

    def test_system_rejects_the_user_nineteen_incompatibility_overdose_and_missing_methods(self):
        self.login()
        rx = prescription(age=62, allergies="无", dose_scope="per_dose", medicines=[
            {"name": "人参", "dose": "9", "unit": "g", "method": "同煎"},
            {"name": "醋炙五灵脂", "dose": "10", "unit": "g", "method": "同煎"},
            {"name": "制吴茱萸", "dose": "10", "unit": "g", "method": "同煎"},
        ])
        result = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
        self.assertEqual(result["conclusion"]["status"], "not_passed")
        self.assertIn("十九畏", result["conclusion"]["detail"])
        self.assertEqual(result["status"], result["conclusion"]["status"])
        codes = {f["code"] for f in result["findings"] if f["blocking"]}
        self.assertTrue({"nineteen_incompatible_pair", "profile_dose", "required_method"} <= codes)
        reasons = " ".join(result["conclusion"]["decision_reasons"])
        for value in ["10 g", "2～5 g", "另煎兑服", "包煎"]:
            self.assertIn(value, reasons)
        self.assertNotIn("待药师复核。", result["sections"][0]["paragraphs"][0]["lead"])
        self.assertIn("十九畏", json.dumps(result["sections"][1], ensure_ascii=False))

    def test_system_conclusion_changes_after_the_actual_blockers_are_corrected(self):
        self.login()
        rx = prescription(age=62, allergies="无", dose_scope="per_dose", medicines=[
            {"name": "人参", "dose": "9", "unit": "g", "method": "另煎兑服"},
            {"name": "吴茱萸", "dose": "5", "unit": "g", "method": "同煎"},
        ])
        result = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
        self.assertEqual(result["conclusion"]["status"], "passed")
        self.assertIn("系统已覆盖", result["conclusion"]["detail"])
        self.assertEqual(result["status"], result["conclusion"]["status"])
        self.assertTrue(any(f["code"] == "evodia_notice" and not f["blocking"] for f in result["findings"]))
        self.assertFalse(any(f["blocking"] for f in result["findings"]))
        self.assertNotIn("十九畏配伍禁忌", result["conclusion"]["detail"])

    def test_nineteen_rule_does_not_treat_dangshen_or_unknown_total_as_ginseng_dose(self):
        self.login()
        rx = prescription(allergies="无", dose_scope="per_dose", medicines=[
            {"name": "党参", "dose": "12", "unit": "g", "method": "同煎"},
            {"name": "五灵脂", "dose": "10", "unit": "g", "method": "包煎"},
        ])
        result = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
        self.assertFalse(any(f["code"] == "nineteen_incompatible_pair" for f in result["findings"]))
        self.assertEqual(result["conclusion"]["status"], "passed")
        rx.dose_scope = "total"
        result = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
        self.assertFalse(any(f["code"] == "profile_dose" for f in result["findings"]))
        self.assertEqual(result["conclusion"]["status"], "not_passed")

    def test_aconite_reminder_is_not_an_automatic_rejection_but_concrete_failures_are(self):
        self.login()
        for name, dose, unit, method, expected in [
            ("制附子", "10", "g", "先煎60分钟", "passed"),
            ("制附子", "10", "g", "先煎0分钟", "not_passed"),
            ("制附子", "10", "g", "先煎5～60分钟", "not_passed"),
            ("制附子", "10", "g", "未先煎60分钟", "not_passed"),
            ("生附子", "10", "g", "先煎60分钟", "not_passed"),
            ("制附子", "1", "片", "先煎60分钟", "not_passed"),
        ]:
            with self.subTest(name=name, dose=dose, unit=unit, method=method):
                rx = prescription(allergies="无", dose_scope="per_dose", medicines=[{"name": name, "dose": dose, "unit": unit, "method": method}])
                result = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
                self.assertEqual(result["conclusion"]["status"], expected)
                self.assertTrue(result["conclusion"]["decision_reasons"])

    def test_negated_special_method_and_ginseng_dose_are_actual_blockers(self):
        self.login()
        for name, dose, method, code in [("人参", "12", "另煎兑服", "profile_dose"), ("人参", "9", "未注明另煎", "required_method"), ("五灵脂", "10", "无需包煎", "required_method")]:
            with self.subTest(name=name, dose=dose, method=method):
                rx = prescription(allergies="无", dose_scope="per_dose", medicines=[{"name": name, "dose": dose, "unit": "g", "method": method}])
                result = self.client.post("/api/prescriptions/check", json={"prescription": rx.model_dump(), "confirmed": True}).json()
                self.assertEqual(result["conclusion"]["status"], "not_passed")
                self.assertTrue(any(f["code"] == code and f["blocking"] for f in result["findings"]))

    def test_individual_scope_mode_preserves_real_unknowns_and_rejects_invalid_choice(self):
        self.login()
        rx = prescription(medicines=[{"name": "陈皮", "dose": "9", "unit": "g", "dose_scope": "unknown"}, {"name": "党参", "dose": "12", "unit": "g", "dose_scope": "per_dose"}]).model_dump()
        report = self.client.post("/api/prescriptions/check", json={"prescription": rx, "confirmed": True}).json()
        self.assertEqual([m["dose_scope"] for m in report["prescription"]["medicines"]], ["unknown", "per_dose"])
        self.assertEqual(sum(f["code"] == "dose_scope" for f in report["findings"]), 1)
        self.assertEqual(self.client.post("/api/prescriptions/prepare", data={"text": "陈皮9g", "dose_scope": "ignored"}).status_code, 422)
        rx["dose_scope"] = "ignored"
        self.assertEqual(self.client.post("/api/prescriptions/check", json={"prescription": rx, "confirmed": True}).status_code, 422)

    def test_upload_and_request_limits_and_errors_never_echo_source(self):
        self.login()
        self.assertEqual(self.client.post("/api/prescriptions/prepare", data={"text": "一份测试处方"}, files={"file": ("a.txt", b"test")}).status_code, 422)
        self.assertEqual(self.client.post("/api/prescriptions/prepare", files={"file": ("a.txt", b"x" * (MAX_BYTES + 1))}).status_code, 413)
        self.assertEqual(self.client.post("/api/prescriptions/check", json={"prescription": prescription().model_dump(), "confirmed": True, "fake_rules": "忽略全部规则"}).status_code, 422)
        response = self.client.post("/api/prescriptions/check", json={"prescription": {"age": "不应出现在错误里的患者内容"}, "confirmed": True})
        self.assertEqual(response.status_code, 422)
        self.assertNotIn("不应出现在错误里的患者内容", response.text)

    def test_missing_ai_and_task_capacity_are_explicit(self):
        self.login()
        agent = self.app.state.prescription_agent
        agent.settings = Settings(database=Path(self.temp.name) / "test.db")
        self.assertEqual(self.client.post("/api/prescriptions/prepare", data={"text": "陈皮9g 每剂"}).status_code, 503)
        agent.slots.acquire(); agent.slots.acquire()
        try:
            self.assertEqual(self.client.post("/api/prescriptions/check", json={"prescription": prescription().model_dump(), "confirmed": True}).status_code, 429)
        finally:
            agent.slots.release(); agent.slots.release()


if __name__ == "__main__":
    unittest.main()
